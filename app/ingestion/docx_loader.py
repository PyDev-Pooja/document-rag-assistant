"""DOCX and plain-text extraction."""
from __future__ import annotations

from pathlib import Path

from app.ingestion.models import ExtractedDocument, ExtractedElement


def extract_docx(path: Path) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX processing") from exc

    document = Document(path)
    elements: list[ExtractedElement] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        element_type = "heading" if style.lower().startswith("heading") else "paragraph"
        elements.append(
            ExtractedElement(
                text=text,
                element_type=element_type,
                source=str(path),
                metadata={"style": style},
            )
        )

    for table_index, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
        table_text = "\n".join(rows).strip()
        if table_text:
            elements.append(
                ExtractedElement(
                    text=table_text,
                    element_type="table",
                    source=str(path),
                    metadata={"table_index": table_index},
                )
            )

    return ExtractedDocument(
        document_id=path.stem,
        filename=path.name,
        source_path=path,
        file_type="docx",
        elements=elements,
        metadata={"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
    )


def extract_txt(path: Path) -> ExtractedDocument:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    elements = []
    if text:
        elements.append(ExtractedElement(text=text, element_type="paragraph", source=str(path)))
    return ExtractedDocument(
        document_id=path.stem,
        filename=path.name,
        source_path=path,
        file_type="txt",
        elements=elements,
    )
